"""Excel reader and test data utility for Purchase Order Update module."""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill
from typing import Any, Dict, List, Optional

_MODULE_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CREDENTIALS_PATH = os.path.join(_MODULE_ROOT, "test_data", "credentials.xlsx")
TEST_WORKBOOK_PATH = os.path.join(_MODULE_ROOT, "test_data", "Update-Purchase-Order.xlsx")


def read_credentials(role: str = "Admin", path: Optional[str] = None) -> Dict[str, Any]:
    """Read user credentials from credentials.xlsx by role."""
    file_path = path or CREDENTIALS_PATH
    if not os.path.exists(file_path):
        fallback = os.path.join(
            os.path.dirname(os.path.dirname(_MODULE_ROOT)),
            "swarajya-create",
            "customer-management",
            "test_data",
            "credentials.xlsx",
        )
        if os.path.exists(fallback):
            file_path = fallback
        else:
            return {
                "employee_id": "332",
                "password": "test@1234",
                "auth_code": "111111",
                "role": role,
            }

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb["Credentials"] if "Credentials" in wb.sheetnames else wb.active

    headers = [str(cell.value).strip().lower() if cell.value else "" for cell in sheet[1]]
    role_idx = headers.index("role") if "role" in headers else 0
    emp_idx = headers.index("employee_id") if "employee_id" in headers else 1
    pass_idx = headers.index("password") if "password" in headers else 2
    auth_idx = headers.index("auth_code") if "auth_code" in headers else 3

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        row_role = str(row[role_idx]).strip() if row[role_idx] else ""
        if row_role.lower() == role.lower():
            return {
                "employee_id": str(row[emp_idx]).strip() if row[emp_idx] else "332",
                "password": str(row[pass_idx]).strip() if row[pass_idx] else "test@1234",
                "auth_code": str(row[auth_idx]).strip() if len(row) > auth_idx and row[auth_idx] else "111111",
                "role": row_role,
            }

    return {
        "employee_id": "332",
        "password": "test@1234",
        "auth_code": "111111",
        "role": role,
    }


def load_test_cases(sheet_name: str, file_path: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    Read all test case rows from given sheet in Update-Purchase-Order.xlsx.
    Returns empty list if file or sheet does not exist.
    """
    path = file_path or TEST_WORKBOOK_PATH
    if not os.path.exists(path):
        return []

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return []

    target_sheet = None
    if sheet_name in wb.sheetnames:
        target_sheet = sheet_name
    else:
        for s in wb.sheetnames:
            if sheet_name.lower() in s.lower() or s.lower() in sheet_name.lower():
                target_sheet = s
                break

    if not target_sheet:
        wb.close()
        return []

    ws = wb[target_sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    headers = [str(cell).strip() if cell is not None else f"col_{idx}" for idx, cell in enumerate(rows[0])]
    cases = []
    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        row_dict = dict(zip(headers, row))
        tc_id = str(row_dict.get("Test Case ID", "") or "").strip()
        if not tc_id or tc_id.lower() == "none":
            continue
        row_dict["_row_idx"] = row_idx
        row_dict["_sheet"] = target_sheet
        cases.append(row_dict)

    return cases


def build_automation_id(tc_id: str) -> str:
    """Generate standardized automation test ID."""
    return f"AUTO_{tc_id.upper()}"


def update_test_result(
    tc_id: str,
    status: str = "Passed",
    remarks: str = "",
    duration: Optional[float] = None,
    file_path: Optional[str] = None,
) -> bool:
    """Update execution status and remarks in Update-Purchase-Order.xlsx."""
    path = file_path or TEST_WORKBOOK_PATH
    if not os.path.exists(path):
        return False

    try:
        wb = openpyxl.load_workbook(path)
        updated = False

        for sheet_name in ["Positive_Tests", "Negative_Tests"]:
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]

            # Find columns
            headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
            tc_col = headers.index("test case id") + 1 if "test case id" in headers else 1
            status_col = headers.index("test status") + 1 if "test status" in headers else 9
            remarks_col = headers.index("remarks") + 1 if "remarks" in headers else 11

            for row in range(2, ws.max_row + 1):
                cell_val = str(ws.cell(row=row, column=tc_col).value or "").strip()
                if cell_val.upper() == tc_id.strip().upper():
                    status_cell = ws.cell(row=row, column=status_col)
                    st_lower = status.strip().lower()

                    if st_lower in ["passed", "pass"]:
                        status_cell.value = "Passed"
                        status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                        status_cell.font = Font(name="Segoe UI", size=10, bold=True, color="006100")
                    elif st_lower in ["failed", "fail"]:
                        status_cell.value = "Failed"
                        status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                        status_cell.font = Font(name="Segoe UI", size=10, bold=True, color="9C0006")
                    elif st_lower in ["skipped", "skip"]:
                        status_cell.value = "Skipped"
                        status_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                        status_cell.font = Font(name="Segoe UI", size=10, bold=True, color="9C6500")
                    else:
                        status_cell.value = status.capitalize()

                    if remarks:
                        remark_text = remarks
                        if duration:
                            remark_text = f"{remarks} ({duration:.2f}s)"
                        ws.cell(row=row, column=remarks_col).value = remark_text
                    updated = True
                    break

        if updated:
            wb.save(path)
        return updated
    except Exception:
        return False
