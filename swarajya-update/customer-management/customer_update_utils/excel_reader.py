import glob
import os
import re
import sys
from datetime import datetime
from typing import Dict, List, Optional

_DIR = os.path.dirname(os.path.abspath(__file__))
_MODULE_ROOT = os.path.dirname(_DIR)
_UPDATE_ROOT = os.path.dirname(_MODULE_ROOT)
_WORKSPACE_ROOT = os.path.dirname(_UPDATE_ROOT)

for _p in (_MODULE_ROOT, _UPDATE_ROOT, _WORKSPACE_ROOT):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import openpyxl
from openpyxl.styles import Font, PatternFill

from shared.utils.logger import get_logger

log = get_logger("customer_excel_reader")
TEST_DATA_DIR = os.path.join(_MODULE_ROOT, "test_data")
DEFAULT_EXCEL_PATH = os.path.join(TEST_DATA_DIR, "Update-Customer-Management.xlsx")
CREDENTIALS_PATH = os.path.join(TEST_DATA_DIR, "credentials.xlsx")

_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def build_automation_id(tc_id: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", tc_id.strip()).strip("_")
    cleaned = re.sub(r"^TC_", "", cleaned, flags=re.IGNORECASE)
    return f"AUT_{cleaned.upper()}"


def resolve_excel_path() -> Optional[str]:
    """Find test case workbook path from env, explicit default, or test_data discovery."""
    custom_path = os.environ.get("CUSTOMER_UPDATE_WORKBOOK")
    if custom_path and os.path.exists(custom_path):
        return custom_path

    if os.path.exists(DEFAULT_EXCEL_PATH):
        return DEFAULT_EXCEL_PATH

    # Search for any candidate Excel workbook in test_data
    if os.path.exists(TEST_DATA_DIR):
        candidates = [
            f for f in glob.glob(os.path.join(TEST_DATA_DIR, "*.xlsx"))
            if not os.path.basename(f).startswith("~$") and "credential" not in os.path.basename(f).lower()
        ]
        if candidates:
            return candidates[0]

    return None


def load_test_cases(sheet_name: str) -> List[Dict]:
    """
    Read all test case rows from given sheet in Excel workbook.
    Returns empty list gracefully if workbook or sheet is not yet present.
    """
    file_path = resolve_excel_path()
    if not file_path or not os.path.exists(file_path):
        log.info(f"No test case Excel file found in {TEST_DATA_DIR}. Awaiting user-added workbook.")
        return []

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    except Exception as exc:
        log.warning(f"Could not open workbook {file_path}: {exc}")
        return []

    target_sheet = None
    if sheet_name in wb.sheetnames:
        target_sheet = sheet_name
    else:
        # Case-insensitive or substring match
        for s in wb.sheetnames:
            if sheet_name.lower() in s.lower() or s.lower() in sheet_name.lower():
                target_sheet = s
                break

    if not target_sheet:
        log.info(f"Sheet '{sheet_name}' not found in {file_path} (Available: {wb.sheetnames})")
        wb.close()
        return []

    ws = wb[target_sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    headers = [str(cell).strip() if cell is not None else f"col_{idx}" for idx, cell in enumerate(rows[0])]
    cases = []
    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        row_dict = dict(zip(headers, row))
        tc_id = str(row_dict.get("Test Case ID", "") or "").strip()
        if not tc_id or tc_id.lower() == "none":
            continue
        row_dict["_row_idx"] = row_idx
        row_dict["_sheet"] = target_sheet
        cases.append(row_dict)

    log.info(f"Loaded {len(cases)} test cases from '{target_sheet}' ({os.path.basename(file_path)})")
    return cases


def read_credentials(role: str = "Admin") -> Dict[str, str]:
    """Read user credentials from credentials.xlsx."""
    candidates = [
        CREDENTIALS_PATH,
        os.path.join(_MODULE_ROOT, "..", "employee-management", "test_data", "credentials.xlsx"),
        os.path.join(_MODULE_ROOT, "..", "..", "swarajya-create", "customer-management", "test_data", "credentials.xlsx"),
    ]

    target_path = next((p for p in candidates if os.path.exists(p)), None)
    if not target_path:
        return {"employee_id": "332", "password": "test@1234", "auth_code": "111111"}

    try:
        wb = openpyxl.load_workbook(target_path, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not rows or len(rows) < 2:
            return {"employee_id": "332", "password": "test@1234", "auth_code": "111111"}

        headers = [str(h or "").strip().lower() for h in rows[0]]
        role_idx = headers.index("role") if "role" in headers else 0
        id_idx = next((i for i, h in enumerate(headers) if "id" in h), 1)
        pwd_idx = next((i for i, h in enumerate(headers) if "pass" in h), 2)
        code_idx = next((i for i, h in enumerate(headers) if "code" in h), 3)

        target_role = role.strip().upper()
        fallback_row = None
        for r in rows[1:]:
            if not r or not any(r):
                continue
            if fallback_row is None:
                fallback_row = r
            if str(r[role_idx] or "").strip().upper() == target_role:
                return {
                    "role": str(r[role_idx] or ""),
                    "employee_id": str(r[id_idx] or "").strip(),
                    "password": str(r[pwd_idx] or "").strip(),
                    "auth_code": str(r[code_idx] or "").strip(),
                }

        if fallback_row:
            return {
                "role": str(fallback_row[role_idx] or ""),
                "employee_id": str(fallback_row[id_idx] or "").strip(),
                "password": str(fallback_row[pwd_idx] or "").strip(),
                "auth_code": str(fallback_row[code_idx] or "").strip(),
            }
    except Exception as exc:
        log.warning(f"Could not read credentials from {target_path}: {exc}")

    return {"employee_id": "332", "password": "test@1234", "auth_code": "111111"}


def update_test_result(
    tc_id: str,
    status: str,
    remarks: str = "",
    duration_s: float = 0.0,
    sheet_name: Optional[str] = None,
):
    """Update execution results in the Excel workbook."""
    file_path = resolve_excel_path()
    if not file_path or not os.path.exists(file_path):
        return

    try:
        wb = openpyxl.load_workbook(file_path)
        sheets_to_search = [wb[sheet_name]] if (sheet_name and sheet_name in wb.sheetnames) else wb.worksheets

        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        clean_status = status.upper().strip()
        fill = _GREEN if clean_status == "PASS" else _RED if clean_status == "FAIL" else _YELLOW

        for ws in sheets_to_search:
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
            if "Test Case ID" not in headers:
                continue

            tc_col = headers.index("Test Case ID") + 1
            status_col = headers.index("Test Status") + 1 if "Test Status" in headers else None
            auto_id_col = headers.index("Auto Script ID") + 1 if "Auto Script ID" in headers else None
            auto_stat_col = headers.index("Automation Status") + 1 if "Automation Status" in headers else None
            time_col = headers.index("Execution_Timestamp") + 1 if "Execution_Timestamp" in headers else None
            remarks_col = headers.index("Remarks") + 1 if "Remarks" in headers else None

            for row in range(2, ws.max_row + 1):
                val = ws.cell(row=row, column=tc_col).value
                if val and str(val).strip() == tc_id.strip():
                    if status_col:
                        cell = ws.cell(row=row, column=status_col, value=clean_status)
                        cell.fill = fill
                        cell.font = Font(bold=True)
                    if auto_id_col:
                        ws.cell(row=row, column=auto_id_col, value=build_automation_id(tc_id))
                    if auto_stat_col:
                        ws.cell(row=row, column=auto_stat_col, value="Automated")
                    if time_col:
                        ws.cell(row=row, column=time_col, value=now_str)
                    if remarks_col and remarks:
                        rem_val = f"{remarks[:250]} ({duration_s:.1f}s)" if duration_s else remarks[:250]
                        ws.cell(row=row, column=remarks_col, value=rem_val)
                    break

        wb.save(file_path)
        wb.close()
    except Exception as exc:
        log.warning(f"Could not update Excel for {tc_id}: {exc}")
