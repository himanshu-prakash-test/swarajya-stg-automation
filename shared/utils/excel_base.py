import os
import re
from datetime import datetime
from typing import Any, Dict, List, Optional
import openpyxl
from shared.utils.logger import get_logger

log = get_logger("SharedExcelBase")

# Resolve shared/test_data directory dynamically
_UTILS_DIR = os.path.dirname(os.path.abspath(__file__))
_SHARED_DIR = os.path.dirname(_UTILS_DIR)
TEST_DATA_DIR = os.path.join(_SHARED_DIR, "test_data")
TEST_DATA_FILE = os.path.join(TEST_DATA_DIR, "Swarajya-test-data.xlsx")
TEST_CASES_FILE = os.path.join(TEST_DATA_DIR, "Swarajya-test-cases.xlsx")


def get_test_data_path(filename: str = "Swarajya-test-data.xlsx") -> str:
    """Return absolute path to a file inside shared/test_data."""
    path = os.path.join(TEST_DATA_DIR, filename)
    if not os.path.exists(path):
        raise FileNotFoundError(f"Requested test data file not found at: {path}")
    return path


def get_test_cases_path(filename: str = "Swarajya-test-cases.xlsx") -> str:
    """Return absolute path to test cases matrix in shared/test_data."""
    return get_test_data_path(filename)


def read_credentials(role: str = "Admin", path: Optional[str] = None) -> Dict[str, str]:
    """
    Strictly fetch user credentials by role from the test data Excel workbook.
    No hardcoded credentials are used.
    """
    file_path = path or TEST_DATA_FILE
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Credentials workbook not found at: {file_path}")

    wb = openpyxl.load_workbook(file_path, data_only=True)
    target_role = role.strip().lower()

    # 1. Search in dedicated role authentication sheets
    for sheet_name in ["loginhr_admin", "loginemp-manager"]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
        role_idx = headers.index("role") if "role" in headers else 0
        id_idx = next((i for i, h in enumerate(headers) if "id" in h), 1)
        pwd_idx = next((i for i, h in enumerate(headers) if "pass" in h), 2)
        auth_idx = next((i for i, h in enumerate(headers) if "code" in h or "auth" in h), 3)

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[role_idx]:
                continue
            row_role = str(row[role_idx]).strip()
            if row_role.lower() == target_role:
                emp_id = row[id_idx]
                auth_code = row[auth_idx]
                emp_id_str = str(int(emp_id)) if isinstance(emp_id, float) else str(emp_id).strip()
                auth_code_str = str(int(auth_code)) if isinstance(auth_code, float) else str(auth_code).strip()
                wb.close()
                return {
                    "role": row_role,
                    "employee_id": emp_id_str,
                    "password": str(row[pwd_idx]).strip(),
                    "auth_code": auth_code_str,
                }

    # 2. Search for embedded credentials in module sheets header comment
    cred_pattern = re.compile(
        r"User\s*ID\s*:\s*(\S+)\s*\|\s*Password\s*:\s*(\S+)\s*\|\s*Auth\s*Code\s*:\s*(\S+)",
        re.IGNORECASE,
    )
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        cell_val = str(ws.cell(row=1, column=1).value or "").strip()
        match = cred_pattern.search(cell_val)
        if match:
            emp_id, pwd, auth_code = match.groups()
            wb.close()
            return {
                "role": role,
                "employee_id": emp_id.strip(),
                "password": pwd.strip(),
                "auth_code": auth_code.strip(),
            }

    wb.close()
    raise ValueError(f"No credentials found for role '{role}' in {file_path}")


def load_test_data(sheet_name: str, file_path: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    Load test input data records from the test data workbook.
    Automatically handles metadata comments on row 1 and extracts column headers.
    """
    path = file_path or TEST_DATA_FILE
    if not os.path.exists(path):
        raise FileNotFoundError(f"Test data workbook not found at: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    matched_sheet = None
    target_clean = sheet_name.lower().replace("-", "").replace("_", "")
    for name in wb.sheetnames:
        if name.lower().replace("-", "").replace("_", "") == target_clean:
            matched_sheet = name
            break

    if not matched_sheet:
        wb.close()
        raise ValueError(f"Sheet '{sheet_name}' not found in {path}. Available sheets: {wb.sheetnames}")

    ws = wb[matched_sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    # Check if first row is a metadata header comment (e.g. "User ID: ...")
    if rows[0] and rows[0][0] and "user id:" in str(rows[0][0]).lower():
        header_row = rows[1]
        data_rows = rows[2:]
    else:
        header_row = rows[0]
        data_rows = rows[1:]

    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(header_row)]
    results = []

    for r in data_rows:
        if not r or not any(x is not None for x in r):
            continue
        record = {}
        for h, val in zip(headers, r):
            if isinstance(val, float) and val.is_integer():
                val = str(int(val))
            elif isinstance(val, datetime):
                val = val.strftime("%d-%m-%Y")
            elif val is not None:
                val = str(val).strip()
            record[h] = val
        results.append(record)

    log.info(f"Loaded {len(results)} records from sheet '{matched_sheet}'")
    return results


def load_test_cases(sheet_name: str, file_path: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    Load test case definitions from the test cases workbook (Swarajya-test-cases.xlsx).
    """
    path = file_path or TEST_CASES_FILE
    if not os.path.exists(path):
        raise FileNotFoundError(f"Test cases workbook not found at: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    matched_sheet = None
    target_clean = sheet_name.lower().replace("-", "").replace("_", "")
    for name in wb.sheetnames:
        if name.lower().replace("-", "").replace("_", "") == target_clean:
            matched_sheet = name
            break

    if not matched_sheet:
        wb.close()
        raise ValueError(f"Sheet '{sheet_name}' not found in {path}. Available sheets: {wb.sheetnames}")

    ws = wb[matched_sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    cases = []
    for r in rows[1:]:
        if not r or not any(x is not None for x in r):
            continue
        case_dict = {}
        for h, val in zip(headers, r):
            if isinstance(val, float) and val.is_integer():
                val = str(int(val))
            elif isinstance(val, datetime):
                val = val.strftime("%d-%m-%Y")
            elif val is not None:
                val = str(val).strip()
            case_dict[h] = val
        cases.append(case_dict)

    log.info(f"Loaded {len(cases)} test cases from sheet '{matched_sheet}'")
    return cases


def load_suite_cases(sheet_name: str, flow_filter: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    Unified suite loader:
    1. Loads test cases from Swarajya-test-cases.xlsx for the specified module sheet.
    2. Merges input test data from Swarajya-test-data.xlsx by matching 'Test Case ID'.
    3. Populates 'Test Data' key-value string and individual data fields.
    4. Optionally filters by flow_filter ('positive' or 'negative').
    """
    cases = load_test_cases(sheet_name)
    try:
        data_records = load_test_data(sheet_name)
        data_map = {}
        for rec in data_records:
            tc_id = str(rec.get("Test Case ID", "")).strip()
            if tc_id:
                data_map[tc_id] = rec
    except Exception:
        data_map = {}

    merged = []
    for case in cases:
        tc_id = str(case.get("Test Case ID", "")).strip()
        if not tc_id:
            continue

        # If data is present, merge fields
        if tc_id in data_map:
            rec = data_map[tc_id]
            for k, v in rec.items():
                if k not in case or not case[k]:
                    case[k] = v
            # If 'Test Data' is empty, construct key-value string
            if not case.get("Test Data"):
                case["Test Data"] = "\n".join(
                    f"{k}: {v}" for k, v in rec.items() if k != "Test Case ID" and v is not None
                )

        # Apply flow filter
        if flow_filter:
            flow_upper = flow_filter.upper()
            if "POS" in flow_upper and "_POS_" not in tc_id:
                continue
            if "NEG" in flow_upper and "_NEG_" not in tc_id:
                continue

        merged.append(case)

    return merged


def is_ui_case(tc: Dict[str, Any]) -> bool:
    """Filter out non-UI / API cases."""
    exec_type = str(tc.get("Execution Type", "")).strip().lower()
    return exec_type != "api"


def build_automation_id(prefix: str, tc_id: str) -> str:
    """Generate normalized script ID like AUT_VENDOR_POS_01."""
    clean_id = re.sub(r"[^A-Za-z0-9]+", "_", tc_id.strip()).strip("_")
    clean_id = re.sub(r"^TC_", "", clean_id, flags=re.IGNORECASE)
    return f"AUT_{clean_id.upper()}"


def update_test_result(
    sheet_name: str,
    tc_id: str,
    status: str,
    auto_id: str = "",
    file_path: Optional[str] = None,
):
    """Update execution status and automation script ID in the test cases Excel workbook."""
    path = file_path or TEST_CASES_FILE
    if not os.path.exists(path):
        log.warning(f"File {path} not found for updating test result.")
        return

    wb = openpyxl.load_workbook(path)
    matched_sheet = None
    target_clean = sheet_name.lower().replace("-", "").replace("_", "")
    for name in wb.sheetnames:
        if name.lower().replace("-", "").replace("_", "") == target_clean:
            matched_sheet = name
            break

    if not matched_sheet:
        wb.close()
        log.warning(f"Sheet '{sheet_name}' not found in {path}")
        return

    ws = wb[matched_sheet]
    headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
    tc_col = headers.index("Test Case ID") + 1 if "Test Case ID" in headers else 1
    status_col = headers.index("Test Status") + 1 if "Test Status" in headers else None
    auto_status_col = headers.index("Automation Status") + 1 if "Automation Status" in headers else None
    auto_id_col = headers.index("Auto Script ID") + 1 if "Auto Script ID" in headers else None

    # If Auto Script ID column doesn't exist, append it
    if auto_id and not auto_id_col:
        auto_id_col = len(headers) + 1
        ws.cell(row=1, column=auto_id_col, value="Auto Script ID")

    for row in range(2, ws.max_row + 1):
        val = str(ws.cell(row=row, column=tc_col).value or "").strip()
        if val == tc_id:
            if status_col:
                ws.cell(row=row, column=status_col, value=status)
            if auto_status_col:
                ws.cell(row=row, column=auto_status_col, value="Automated")
            if auto_id and auto_id_col:
                ws.cell(row=row, column=auto_id_col, value=auto_id)
            break

    wb.save(path)
    wb.close()
    log.info(f"Updated Excel {tc_id}: Test Status={status}, Auto Script ID={auto_id}")
