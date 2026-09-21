import glob
import os
import time
from datetime import datetime
from typing import Any, Dict, List, Optional
import openpyxl
from shared.utils.logger import get_logger

logger = get_logger("po_excel_reader")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEST_DATA_DIR = os.path.join(ROOT, "test_data")
CREDENTIALS_FILE = os.path.join(TEST_DATA_DIR, "credentials.xlsx")


def get_test_cases_filepath() -> Optional[str]:
    """Find Purchase Order test cases Excel file in test_data directory."""
    if not os.path.exists(TEST_DATA_DIR):
        return None

    # Common exact and pattern matches
    preferred_names = [
        "Create-Purchase-Order-Test-Cases.xlsx",
        "Create-Purchase-Order.xlsx",
        "Create-Purchase-Orders.xlsx",
        "Update-Purchase-Order.xlsx",
        "Purchase-Order.xlsx",
        "Purchase-Orders.xlsx",
        "Swarajya-Create-Purchase-Order.xlsx",
    ]
    for pref in preferred_names:
        full_path = os.path.join(TEST_DATA_DIR, pref)
        if os.path.exists(full_path):
            return full_path

    # Glob matching for purchase/po excel files
    candidates = (
        glob.glob(os.path.join(TEST_DATA_DIR, "*[Cc]reate*[Pp]urchase*.[xX][lL][sS]*"))
        + glob.glob(os.path.join(TEST_DATA_DIR, "*[Pp]urchase*.[xX][lL][sS]*"))
        + glob.glob(os.path.join(TEST_DATA_DIR, "*[pP][oO]*.[xX][lL][sS]*"))
        + glob.glob(os.path.join(TEST_DATA_DIR, "*.xlsx"))
    )

    # Filter out temporary files and credentials
    candidates = [
        c for c in candidates
        if not os.path.basename(c).startswith("~$")
        and "credential" not in os.path.basename(c).lower()
    ]

    if candidates:
        return candidates[0]
    return None


def read_credentials(role: str = "Admin") -> Dict[str, str]:
    """Read login credentials for the given role from credentials.xlsx."""
    if not os.path.exists(CREDENTIALS_FILE):
        return {"role": "Admin", "employee_id": "332", "password": "test@1234", "auth_code": "111111"}

    try:
        wb = openpyxl.load_workbook(CREDENTIALS_FILE, data_only=True)
        ws = wb["Credentials"] if "Credentials" in wb.sheetnames else wb.active
        headers = [cell.value for cell in ws[1]]

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            row_dict = dict(zip(headers, row))
            row_role = str(row_dict.get("Role", "")).strip().lower()
            if row_role == role.strip().lower() or (
                role.lower() in ("admin", "manager") and row_role in ("admin", "manager", "employee")
            ):
                wb.close()
                return {
                    "role": str(row_dict.get("Role", "")).strip(),
                    "employee_id": str(row_dict.get("Employee_ID", "")).strip(),
                    "password": str(row_dict.get("Password", "")).strip(),
                    "auth_code": str(row_dict.get("Auth_Code", "111111")).strip(),
                }
        wb.close()
    except Exception as exc:
        logger.warning("Error reading credentials: %s", exc)

    return {"role": "Admin", "employee_id": "332", "password": "test@1234", "auth_code": "111111"}


def read_test_cases(sheet_name: str) -> List[Dict[str, Any]]:
    """Read test case rows from specified sheet in the purchase order Excel workbook."""
    file_path = get_test_cases_filepath()
    if not file_path or not os.path.exists(file_path):
        logger.warning("No test cases file found in %s yet", TEST_DATA_DIR)
        return []

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as exc:
        logger.warning("Could not open workbook %s: %s", file_path, exc)
        return []

    target_sheet = None
    for name in wb.sheetnames:
        if name.strip().lower() == sheet_name.strip().lower():
            target_sheet = name
            break

    if not target_sheet:
        # Fallback to single active sheet if only one sheet exists
        if len(wb.sheetnames) == 1:
            target_sheet = wb.sheetnames[0]
        else:
            wb.close()
            logger.info("Sheet '%s' not found in %s. Available: %s", sheet_name, file_path, wb.sheetnames)
            return []

    ws = wb[target_sheet]
    headers = [cell.value for cell in ws[1]]

    test_cases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        row_dict = dict(zip(headers, row))
        tc_id = row_dict.get("Test Case ID")
        if tc_id:
            row_dict["_sheet"] = target_sheet
            if "Test Steps" in row_dict and "Steps" not in row_dict:
                row_dict["Steps"] = row_dict["Test Steps"]
            test_cases.append(row_dict)

    wb.close()
    logger.info("Loaded %d test cases from '%s' in %s", len(test_cases), target_sheet, os.path.basename(file_path))
    return test_cases


def is_ui_case(tc: Dict[str, Any]) -> bool:
    """Check if test case is active for automation."""
    auto_status = str(tc.get("Automation Status", "")).strip().lower()
    return auto_status not in ("excluded", "deferred", "deprecated")


def build_automation_id(tc_id: str) -> str:
    """Generate automation script ID from test case ID (e.g. TC_PO_POS_01 -> AUT_PO_POS_01)."""
    return tc_id.replace("TC_", "AUT_")


def update_test_result(
    tc_id: str,
    status: str,
    error_message: Optional[str] = None,
    execution_time: Optional[float] = None,
):
    """Update test case execution status and remarks in the Excel sheet."""
    file_path = get_test_cases_filepath()
    if not file_path or not os.path.exists(file_path):
        return

    try:
        wb = openpyxl.load_workbook(file_path)
        updated = False
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            headers = [cell.value for cell in ws[1]]
            if "Test Case ID" not in headers:
                continue

            tc_id_col = headers.index("Test Case ID") + 1
            status_col = headers.index("Test Status") + 1 if "Test Status" in headers else None
            remarks_col = headers.index("Remarks") + 1 if "Remarks" in headers else None
            auto_id_col = headers.index("Auto Script ID") + 1 if "Auto Script ID" in headers else None

            for row_idx in range(2, ws.max_row + 1):
                val = ws.cell(row=row_idx, column=tc_id_col).value
                if val and str(val).strip() == tc_id:
                    if status_col:
                        ws.cell(row=row_idx, column=status_col, value=status)
                    if auto_id_col:
                        ws.cell(row=row_idx, column=auto_id_col, value=build_automation_id(tc_id))
                    if remarks_col:
                        now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                        remarks_val = f"{now_str} - {status}"
                        if error_message:
                            remarks_val += f": {error_message[:150]}"
                        ws.cell(row=row_idx, column=remarks_col, value=remarks_val)
                    updated = True
                    break

            if updated:
                break

        if updated:
            wb.save(file_path)
            logger.info("Updated Excel result for %s: %s", tc_id, status)
        wb.close()
    except Exception as exc:
        logger.warning("Could not update Excel test result for %s: %s", tc_id, exc)
