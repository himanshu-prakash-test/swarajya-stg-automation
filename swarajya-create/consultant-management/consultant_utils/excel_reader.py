import glob
import os
import time
from datetime import datetime
from typing import Any, Dict, List, Optional

import openpyxl

from shared.utils.logger import get_logger

logger = get_logger("excel_reader")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEST_DATA_DIR = os.path.join(ROOT, "test_data")
CREDENTIALS_FILE = os.path.join(TEST_DATA_DIR, "credentials.xlsx")


def get_test_cases_filepath() -> str:
    """Find Create-Consultant-Management.xlsx in test_data directory."""
    target = os.path.join(TEST_DATA_DIR, "Create-Consultant-Management.xlsx")
    if os.path.exists(target):
        return target
    candidates = glob.glob(os.path.join(TEST_DATA_DIR, "*consultant*.xlsx"))
    if candidates:
        return candidates[0]
    raise FileNotFoundError(f"No consultant test cases Excel file found in: {TEST_DATA_DIR}")


def read_credentials(role: str = "Admin") -> Dict[str, str]:
    """Read login credentials for the given role from credentials.xlsx."""
    if not os.path.exists(CREDENTIALS_FILE):
        raise FileNotFoundError(f"Credentials file not found: {CREDENTIALS_FILE}")

    wb = openpyxl.load_workbook(CREDENTIALS_FILE, data_only=True)
    ws = wb["Credentials"] if "Credentials" in wb.sheetnames else wb.active
    headers = [cell.value for cell in ws[1]]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        row_dict = dict(zip(headers, row))
        row_role = str(row_dict.get("Role", "")).strip().lower()
        if row_role == role.strip().lower() or (role.lower() in ("admin", "manager") and row_role in ("admin", "manager")):
            wb.close()
            return {
                "role": str(row_dict.get("Role", "")).strip(),
                "employee_id": str(row_dict.get("Employee_ID", "")).strip(),
                "password": str(row_dict.get("Password", "")).strip(),
                "auth_code": str(row_dict.get("Auth_Code", "111111")).strip(),
            }

    wb.close()
    return {"role": "Admin", "employee_id": "332", "password": "test@1234", "auth_code": "111111"}


def read_test_cases(sheet_name: str) -> List[Dict[str, Any]]:
    """Read test case rows from specified sheet in Create-Consultant-Management.xlsx."""
    file_path = get_test_cases_filepath()
    wb = openpyxl.load_workbook(file_path, data_only=True)

    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet '{sheet_name}' not found in {file_path}. Available: {wb.sheetnames}")

    ws = wb[sheet_name]
    headers = [cell.value for cell in ws[1]]

    test_cases = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        row_dict = dict(zip(headers, row))
        tc_id = row_dict.get("Test Case ID")
        if tc_id and str(tc_id).strip().startswith("TC_CONSULTANT"):
            row_dict["_sheet"] = sheet_name
            if "Test Steps" in row_dict and "Steps" not in row_dict:
                row_dict["Steps"] = row_dict["Test Steps"]
            test_cases.append(row_dict)

    wb.close()
    logger.info("Loaded %d test cases from '%s'", len(test_cases), sheet_name)
    return test_cases


def is_ui_case(tc: Dict[str, Any]) -> bool:
    """Check if test case execution type is UI."""
    exec_type = str(tc.get("Execution Type", "UI")).strip().upper()
    return "API" not in exec_type


def build_automation_id(tc_id: str) -> str:
    """Generate automation script ID from test case ID (e.g. TC_CONSULTANT_POS_01 -> AUT_CONSULTANT_POS_01)."""
    return tc_id.replace("TC_", "AUT_")


def update_test_result(
    tc_id: str,
    status: str,
    error_message: Optional[str] = None,
    execution_time: Optional[float] = None,
):
    """Write test execution result back to Create-Consultant-Management.xlsx."""
    file_path = get_test_cases_filepath()
    for attempt in range(3):
        try:
            wb = openpyxl.load_workbook(file_path)
            for sheet_name in ["Positive_Tests", "Negative_Tests"]:
                if sheet_name not in wb.sheetnames:
                    continue
                ws = wb[sheet_name]
                headers = [cell.value for cell in ws[1]]
                id_col = headers.index("Test Case ID") + 1 if "Test Case ID" in headers else 1

                for row_idx in range(2, ws.max_row + 1):
                    cell_val = ws.cell(row=row_idx, column=id_col).value
                    if cell_val and str(cell_val).strip() == tc_id.strip():
                        def set_col(col_name, val):
                            if col_name in headers:
                                c = headers.index(col_name) + 1
                                ws.cell(row=row_idx, column=c, value=val)

                        set_col("Status", status)
                        set_col("Test Status", status)
                        set_col("Automation Status", "Automated")
                        set_col("Auto Script ID", build_automation_id(tc_id))
                        set_col("Remarks", error_message or "Execution Completed Successfully")
                        if execution_time is not None:
                            set_col("Execution Time (s)", round(execution_time, 2))
                        wb.save(file_path)
                        wb.close()
                        return
            wb.close()
            return
        except Exception:
            time.sleep(0.5)
