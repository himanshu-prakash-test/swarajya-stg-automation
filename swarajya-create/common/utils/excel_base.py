import os
from typing import Any, Dict, List, Optional
import openpyxl
from shared.utils.logger import get_logger

log = get_logger("SharedExcelBase")


def read_credentials(role: str = "Manager", path: Optional[str] = None) -> Dict[str, str]:
    """Read user credentials by role from credentials.xlsx."""
    if not path or not os.path.exists(path):
        raise FileNotFoundError(f"Credentials file not found at: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    headers = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    role_idx = headers.index("role") if "role" in headers else 0
    id_idx = headers.index("employee_id") if "employee_id" in headers else (headers.index("email") if "email" in headers else 1)
    pwd_idx = headers.index("password") if "password" in headers else 2
    auth_idx = headers.index("auth_code") if "auth_code" in headers else -1

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and len(row) > max(role_idx, id_idx, pwd_idx):
            row_role = str(row[role_idx]).strip() if row[role_idx] else ""
            if row_role.lower() == role.lower():
                return {
                    "role": row_role,
                    "employee_id": str(row[id_idx]).strip(),
                    "password": str(row[pwd_idx]).strip(),
                    "auth_code": str(row[auth_idx]).strip() if auth_idx >= 0 and row[auth_idx] else "111111",
                }

    # Fallback to first row
    first_row = list(ws.iter_rows(min_row=2, values_only=True))[0]
    return {
        "role": role,
        "employee_id": str(first_row[id_idx]).strip(),
        "password": str(first_row[pwd_idx]).strip(),
        "auth_code": "111111",
    }


def is_ui_case(tc: Dict[str, Any]) -> bool:
    """Filter out non-UI / API cases."""
    exec_type = str(tc.get("Execution Type", "")).strip().lower()
    return exec_type != "api"


def build_automation_id(prefix: str, tc_id: str) -> str:
    """Generate normalized script ID like AUT_VENDOR_POS_01 or AUT_EMP_POS_01."""
    clean_id = tc_id.replace("TC_", "").replace("TC", "")
    return f"AUT_{clean_id}"


def update_test_result(
    file_path: str,
    sheet_name: str,
    tc_id: str,
    status: str,
    auto_id: str,
):
    """Update execution status, auto script ID, and automated flag in Excel."""
    if not os.path.exists(file_path):
        return

    wb = openpyxl.load_workbook(file_path)
    if sheet_name not in wb.sheetnames:
        return
    ws = wb[sheet_name]

    headers = [str(cell.value).strip() if cell.value else "" for cell in ws[1]]
    tc_col = headers.index("Test Case ID") + 1 if "Test Case ID" in headers else 1
    status_col = headers.index("Test Status") + 1 if "Test Status" in headers else None
    auto_status_col = headers.index("Automation Status") + 1 if "Automation Status" in headers else None
    auto_id_col = headers.index("Auto Script ID") + 1 if "Auto Script ID" in headers else None

    for row in range(2, ws.max_row + 1):
        val = str(ws.cell(row=row, column=tc_col).value or "").strip()
        if val == tc_id:
            if status_col:
                ws.cell(row=row, column=status_col, value=status)
            if auto_status_col:
                ws.cell(row=row, column=auto_status_col, value="Automated")
            if auto_id_col:
                ws.cell(row=row, column=auto_id_col, value=auto_id)
            break

    wb.save(file_path)
    wb.close()
    log.info(f"Updated Excel {tc_id}: Test Status={status}, Auto Script ID={auto_id}")
